"""EOG global gas flare catalogue.

The catalogue is published as one workbook per year with three flare sheets:
upstream, oil downstream and gas downstream. Each row is a flare site with a
coordinate, an average temperature, an ellipticity, a detection frequency and a
flared volume estimate.

Temperature and detection frequency are the two fields that matter most here.
They are the hand set discriminators in the published flare survey method, and
KAALCHAKRA learns what that method sets by threshold.
"""

from pathlib import Path
from typing import Final

import openpyxl

FLARE_SHEETS: Final[tuple[str, ...]] = (
    "flare upstream",
    "flare oil downstream",
    "flare gas downstream",
)

# Column headers vary slightly by year, so they are matched on a prefix rather
# than on an exact string. The year is embedded in several of them.
FIELD_PREFIXES: Final[dict[str, str]] = {
    "country": "country",
    "iso": "iso",
    "flare_id": "id",
    "latitude": "latitude",
    "longitude": "longitude",
    "bcm": "bcm",
    "avg_temp_k": "avg temp",
    "ellipticity": "ellipticity",
    "detection_freq": "detection freq",
    "clear_obs": "clear obs",
    "flare_type": "type",
}


def _value(record: tuple[object, ...], columns: dict[str, int], field: str) -> object:
    """Read one field from a row, tolerating a sheet that lacks the column."""
    index = columns.get(field)
    if index is None or index >= len(record):
        return None
    return record[index]


def _resolve_columns(header: tuple[object, ...]) -> dict[str, int]:
    """Map field names onto column indexes for one sheet's header row."""
    lowered = [str(cell).strip().lower() if cell is not None else "" for cell in header]
    resolved: dict[str, int] = {}
    for field, prefix in FIELD_PREFIXES.items():
        for index, name in enumerate(lowered):
            if name.startswith(prefix):
                resolved[field] = index
                break
    return resolved


def load_year(path: Path, year: int, iso_filter: str = "IND") -> list[dict[str, object]]:
    """Read one annual workbook and return flare rows for one country.

    iso_filter is applied on the catalogue's own ISO code column rather than on a
    bounding box, because a bounding box over India also captures flares in
    Pakistan, Bangladesh and Myanmar that are not this project's subject.
    """
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows: list[dict[str, object]] = []

    for sheet_name in FLARE_SHEETS:
        if sheet_name not in workbook.sheetnames:
            continue
        sheet = workbook[sheet_name]
        iterator = sheet.iter_rows(values_only=True)
        try:
            header = next(iterator)
        except StopIteration:
            continue
        columns = _resolve_columns(header)
        if "latitude" not in columns or "longitude" not in columns:
            continue

        for record in iterator:
            if record is None or columns["latitude"] >= len(record):
                continue
            iso = record[columns["iso"]] if "iso" in columns else None
            if iso is None or str(iso).strip().upper() != iso_filter:
                continue
            latitude = record[columns["latitude"]]
            longitude = record[columns["longitude"]]
            if latitude is None or longitude is None:
                continue

            rows.append(
                {
                    "catalogue_year": year,
                    "sheet": sheet_name,
                    "country": _value(record, columns, "country"),
                    "iso": str(iso).strip().upper(),
                    "flare_id": _value(record, columns, "flare_id"),
                    "latitude": float(latitude),
                    "longitude": float(longitude),
                    "bcm": _value(record, columns, "bcm"),
                    "avg_temp_k": _value(record, columns, "avg_temp_k"),
                    "ellipticity": _value(record, columns, "ellipticity"),
                    "detection_freq": _value(record, columns, "detection_freq"),
                    "clear_obs": _value(record, columns, "clear_obs"),
                    "flare_type": _value(record, columns, "flare_type"),
                }
            )
    workbook.close()
    return rows
